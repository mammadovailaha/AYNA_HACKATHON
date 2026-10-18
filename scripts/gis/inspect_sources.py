"""Excel mənbələrini və problemli seqmentləri yoxlamaq üçün köməkçi skript (yalnız oxuyur)."""
import json
import sys
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

import openpyxl
import xlrd
from pyproj import Geod

sys.stdout.reconfigure(encoding="utf-8")
ROOT = Path(__file__).resolve().parents[2]
RAW = ROOT / "dist" / "raw"
KML = "{http://www.opengis.net/kml/2.2}"
GEOD = Geod(ellps="WGS84")

print("== mikromobillik_teklif.xlsx")
wb = openpyxl.load_workbook(RAW / "mikromobillik_teklif.xlsx", data_only=True)
for ws in wb.worksheets:
    print(f"-- sheet {ws.title} {ws.max_row}x{ws.max_column}")
    for row in ws.iter_rows(values_only=True):
        if any(v is not None for v in row):
            print("  ", [v for v in row if v is not None])

print("\n== 003_1.xls")
book = xlrd.open_workbook(RAW / "003_1.xls")
for sh in book.sheets():
    print(f"-- sheet {sh.name} {sh.nrows}x{sh.ncols}")
    for r in range(sh.nrows):
        vals = [v for v in sh.row_values(r) if v not in ("", None)]
        if vals:
            print("  ", vals[:12])

print("\n== problemli placemark-lar")
for kmz in ("movcud-zolaqlar-velo.kmz", "mpv-v-mps.-son.kmz"):
    with zipfile.ZipFile(RAW / kmz) as z:
        root = ET.fromstring(z.read(next(n for n in z.namelist() if n.endswith(".kml"))))
    for pm in root.iter(f"{KML}Placemark"):
        name = (pm.findtext(f"{KML}name") or "").strip()
        for ls in pm.iter(f"{KML}LineString"):
            coords = [tuple(map(float, c.split(","))) for c in (ls.findtext(f"{KML}coordinates") or "").split()]
            if name in ("8 Noyabr pr.", "Mətbuat pr.") or len(coords) < 2:
                jumps = []
                for a, b in zip(coords, coords[1:]):
                    _, _, d = GEOD.inv(a[0], a[1], b[0], b[1])
                    jumps.append(round(d))
                desc = (pm.findtext(f"{KML}description") or "")[:120]
                print(kmz, "|", name, "| nöqtə:", len(coords), "| ən böyük addım m:", max(jumps) if jumps else None,
                      "| uzunluq km:", round(sum(jumps) / 1000, 3), "| ilk:", coords[0][:2] if coords else None, "| son:", coords[-1][:2] if coords else None, "| desc:", desc)
